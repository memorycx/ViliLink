import os
from datetime import datetime
from common.Config import config
from docx import Document

from utils.docx_utils import generate_docx_report

root_path = config.REPORT_PATH / "运维"


def _clean_parse_situation(key_text):
    """
    专门解析：'情况记录：\n管网巡查\n异常：暂无'
    返回干净的字典：{'情况记录': '管网巡查', '异常': '暂无'}
    """
    result = {}

    # 按换行切分
    lines = [line.strip() for line in key_text.splitlines() if line.strip()]

    # 遍历每一行提取内容
    for line in lines:
        if line.startswith("情况记录"):
            # 下一行就是内容
            idx = lines.index(line)
            if idx + 1 < len(lines):
                result["情况记录"] = lines[idx + 1]

        elif line.startswith("异常"):
            # 提取异常值
            result["异常"] = line.replace("异常：", "").strip()

    return result


def read_docx_table(docx_path):
    """读取docx里的表格数据"""
    doc = Document(docx_path)
    table_data = {}

    # 遍历所有表格
    for table in doc.tables:
        # 遍历表格的每一行
        for row in table.rows:
            # 遍历行里的每一个单元格
            cells = row.cells
            if len(cells) >= 2:
                # 处理像 "乡镇: 芝英镇" 这种格式
                key = cells[0].text.strip().replace(":", "").strip()
                value = cells[1].text.strip()
                table_data[key] = value
            if len(cells) >= 4:
                # 处理 "设施名称: 黄店村  巡查人员: 常鑫" 这种格式
                key1 = cells[0].text.strip().replace(":", "").strip()
                value1 = cells[1].text.strip()
                key2 = cells[2].text.strip().replace(":", "").strip()
                value2 = cells[3].text.strip()
                table_data[key1] = value1
                table_data[key2] = value2

    for key in list(table_data.keys()):
        if "情况记录" in key:
            tem = _clean_parse_situation(key)
            table_data.pop(key)
            table_data.update(tem)

    return table_data


def _extract_images_from_docx(docx_path, output_dir, user):
    """从docx中提取所有图片，保存到指定目录，返回图片路径列表"""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    doc = Document(docx_path)
    image_paths = []

    # 遍历文档中的所有内嵌图片
    imag_id = 0
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            imag_id += 1
            img = rel.target_part
            img_filename = os.path.basename(rel.target_ref)
            name_without_ext, ext = os.path.splitext(img_filename)
            new_filename = f"{user}_{imag_id}{ext}"
            img_path = os.path.join(output_dir, new_filename)
            with open(img_path, "wb") as f:
                f.write(img.blob)
            image_paths.append({
                "image_path": img_path
            })

    return image_paths


def _merge_multiple_dicts(*dicts, sep=" / "):
    """
    合并多个报表字典：
    - 所有字典中该键的值相同 → 保留一个
    - 值不同 → 去重、过滤空值后，用 sep 拼接
    - 跳过空键/空值
    """
    merged = {}

    all_keys = set()
    for d in dicts:
        if d:
            all_keys.update(d.keys())

    for key in all_keys:
        if not key:
            continue

        values = []
        for d in dicts:
            val = str(d.get(key, "")).strip()
            if val:
                values.append(val)

        if not values:
            continue

        unique_values = list(dict.fromkeys(values))
        if len(unique_values) == 1:
            merged[key] = unique_values[0]
        else:
            merged[key] = sep.join(unique_values)

    return merged


def merge_file():
    """
    唯一向外暴露的方法，对同一日同一地点的文件进行合并处理
    :return:
    """
    if not root_path.exists() or not root_path.is_dir():
        print(f"错误：报表路径不存在或无效: {root_path}")
        return

    for town_village_folder in os.listdir(root_path):

        folder_path = os.path.join(root_path, town_village_folder)

        try:
            if len(os.listdir(folder_path)) == 1:
                continue
        except (OSError, PermissionError) as e:
            print(f"警告：无法读取目录 {folder_path}: {e}")
            continue

        all_data = []
        image_paths = []
        today = datetime.now().strftime("%Y-%m-%d")
        for data_docx in os.listdir(folder_path):

            if today in data_docx:
                try:
                    tem = read_docx_table(os.path.join(folder_path, data_docx))
                    parts = data_docx.split("_")
                    if len(parts) >= 2:
                        user = parts[1].split(".")[0]
                    else:
                        user = "unknown"
                    all_data.append(tem)

                    image_paths.extend(_extract_images_from_docx(os.path.join(folder_path, data_docx),
                                                                 os.path.join(root_path, "merge"),
                                                                 user))
                except Exception as e:
                    print(f"警告：处理文件 {data_docx} 时出错: {e}")
                    continue

        if len(all_data) >= 2:
            try:
                end_data = _merge_multiple_dicts(all_data[0], all_data[1], sep="/")

                required_keys = ["乡镇：", "日期：", "设施名称：", "巡查人员：", "处理工艺：", "设施编号：", "情况记录",
                                 "异常"]
                if not all(key in end_data for key in required_keys):
                    print(f"警告：合并后的数据缺少必要字段")
                    continue

                report_data = {
                    "town": end_data["乡镇："],
                    "date": end_data["日期："],
                    "facility_name": end_data["设施名称："],
                    "inspector": end_data["巡查人员："],
                    "process": end_data["处理工艺："],
                    "facility_code": end_data["设施编号："],
                    "record": end_data["情况记录"],
                    "abnormal_status": end_data["异常"]
                }

                tem = end_data['巡查人员：'].replace("/", " ")
                generate_docx_report(
                    os.path.join(folder_path, f"{tem}_{today}.docx"),
                    report_data,
                    image_paths
                )
            except Exception as e:
                print(f"警告：生成合并报表时出错: {e}")
                continue

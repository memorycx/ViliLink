"""合并同一天、同一地点的多份运维报表。"""

from __future__ import annotations

import os
import re
import tempfile
import traceback
from datetime import date
from pathlib import Path
from typing import Callable, Iterable, Mapping
from uuid import uuid4

from docx import Document

from common.config import Settings
from common.models import ABNORMAL_NONE, ReportData
from utils.docx_utils import DEFAULT_REPORT_TITLE, generate_report
from utils.paths import safe_filename, unique_path

# docx 表格中的中文标签 -> ReportData 字段
_FIELD_MAP = {
    "乡镇": "town",
    "日期": "date",
    "设施名称": "facility_name",
    "巡查人员": "inspector",
    "处理工艺": "process",
    "设施编号": "facility_code",
    "情况记录": "record",
    "异常": "abnormal_status",
}

LogFunc = Callable[..., None]


class ReportMerger:
    def __init__(self, settings: Settings, log_func: LogFunc):
        self._settings = settings
        self._log = log_func

    def merge_today(self) -> int:
        """扫描“运维”目录，合并当天同一地点的多份报表。"""
        day_text = date.today().strftime("%Y-%m-%d")
        source_root = self._settings.report_path / "运维"
        if not source_root.is_dir():
            self._log(f"未找到运维报表目录: {source_root}", "WARNING")
            return 0

        merged_count = 0
        for folder in sorted(path for path in source_root.iterdir() if path.is_dir()):
            documents = [
                path
                for path in sorted(folder.glob("*.docx"))
                if day_text in path.name and not path.name.startswith("~$")
            ]
            if len(documents) < 2:
                continue

            # 只合并能识别出报表字段的文件；读取失败的文件保留不动。
            parsed: list[tuple[Path, dict[str, str]]] = []
            for document in documents:
                try:
                    record = read_report_docx(document)
                except Exception as exc:
                    self._log(f"读取报表失败 {document.name}: {exc}", "WARNING")
                    continue
                if not {"town", "date", "inspector"}.issubset(record):
                    self._log(f"{document.name}: 不像标准报表，已跳过", "WARNING")
                    continue
                parsed.append((document, record))

            if len(parsed) < 2:
                continue

            report_data = merge_records(record for _, record in parsed)
            if report_data is None:
                self._log(f"{folder.name}: 合并所需字段不完整，已跳过", "WARNING")
                continue

            try:
                merged_count += self._write_merged_report(
                    folder,
                    [document for document, _ in parsed],
                    report_data,
                )
            except Exception:
                self._log(f"{folder.name}: 生成合并报表失败\n{traceback.format_exc()}", "ERROR")

        self._log(f"合并任务结束，共生成 {merged_count} 份合并报表", "SUCCESS")
        return merged_count

    def _write_merged_report(
        self,
        source_folder: Path,
        documents: list[Path],
        report_data: ReportData,
    ) -> int:
        # 先生成到同目录的临时文件，全部成功后再删除原报表并改名。
        temp_output = source_folder / f".merge_{uuid4().hex}.tmp"

        try:
            with tempfile.TemporaryDirectory(prefix="villink_merge_") as temp_dir:
                image_paths: list[str] = []
                for document in documents:
                    try:
                        image_paths.extend(extract_images(document, Path(temp_dir)))
                    except Exception as exc:
                        self._log(f"提取图片失败 {document.name}: {exc}", "WARNING")

                report_data.image_paths = image_paths
                generate_report(temp_output, report_data, title=DEFAULT_REPORT_TITLE)
        except Exception:
            temp_output.unlink(missing_ok=True)
            raise

        deleted = self._delete_sources(documents)

        target = source_folder / merged_filename(report_data)
        if target.exists():
            # 删除失败时源文件可能还在，避免覆盖它。
            target = unique_path(target)

        try:
            os.replace(temp_output, target)
        except OSError as exc:
            # 原报表已删除，保留临时文件以便手动恢复，不能直接丢弃。
            self._log(f"合并报表改名失败，文件保留在 {temp_output}: {exc}", "ERROR")
            raise

        self._log(
            f"合并完成: {target}（已删除 {deleted}/{len(documents)} 份原报表）",
            "SUCCESS",
        )
        return 1

    def _delete_sources(self, documents: list[Path]) -> int:
        deleted = 0
        for document in documents:
            try:
                document.unlink()
            except OSError as exc:
                self._log(f"删除原报表失败 {document.name}: {exc}", "WARNING")
            else:
                deleted += 1
        return deleted


def read_report_docx(path: str | Path) -> dict[str, str]:
    """读取报表表格，返回字段名 -> 文本。"""
    document = Document(str(path))
    data: dict[str, str] = {}

    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            seen: set[int] = set()
            for cell in row.cells:
                identifier = id(cell._tc)
                if identifier in seen:
                    continue
                seen.add(identifier)
                cells.append(cell.text.strip())

            if len(cells) >= 4:
                _store_pair(data, cells[0], cells[1])
                _store_pair(data, cells[2], cells[3])
            elif len(cells) == 2:
                _store_pair(data, cells[0], cells[1])
            elif len(cells) == 1 and ("情况记录" in cells[0] or "异常" in cells[0]):
                _parse_record_cell(data, cells[0])

    return data


def merge_records(records: Iterable[Mapping[str, str]]) -> ReportData | None:
    """把多份报表的字段合并为一份；不同取值用“ / ”连接。"""
    rows = list(records)
    merged = {
        field: _merge_values(row.get(field, "") for row in rows)
        for field in _FIELD_MAP.values()
        if field != "record"
    }
    merged["record"] = _merge_values(
        (row.get("record", "") for row in rows),
        separator="\n",
    )

    # 合并异常状态时忽略“暂无”，避免出现“暂无 / 管道破裂”这种互相矛盾的展示。
    abnormal_values = [
        str(row.get("abnormal_status", "")).strip()
        for row in rows
        if str(row.get("abnormal_status", "")).strip() not in {"", ABNORMAL_NONE}
    ]
    merged["abnormal_status"] = _merge_values(abnormal_values) if abnormal_values else ABNORMAL_NONE

    required = (
        "town",
        "date",
        "facility_name",
        "inspector",
        "process",
        "facility_code",
        "record",
    )
    if any(not merged.get(field) for field in required):
        return None

    return ReportData(
        town=merged["town"],
        date=merged["date"],
        facility_name=merged["facility_name"],
        inspector=merged["inspector"],
        process=merged["process"],
        facility_code=merged["facility_code"],
        record=merged["record"],
        abnormal_status=merged.get("abnormal_status") or ABNORMAL_NONE,
    )


def merged_filename(report_data: ReportData) -> str:
    """合并报表文件名：日期（不含时间）+ 巡查人员姓名。"""
    match = re.search(r"\d{4}-\d{2}-\d{2}", report_data.date or "")
    date_text = match.group(0) if match else date.today().strftime("%Y-%m-%d")

    names: list[str] = []
    for part in re.split(r"[/、,，;；]", report_data.inspector or ""):
        name = part.strip()
        if name and name not in names:
            names.append(name)

    stem = f"{date_text}_{' '.join(names) or '未知人员'}"
    return safe_filename(stem, f"{date_text}_合并") + ".docx"


def extract_images(docx_path: str | Path, output_dir: str | Path) -> list[str]:
    """提取 docx 内嵌图片，返回临时图片路径（供合并报表重新嵌入）。"""
    docx_path = Path(docx_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    document = Document(str(docx_path))

    paths: list[str] = []
    for index, relation in enumerate(document.part.rels.values(), 1):
        if "image" not in relation.target_ref:
            continue
        part = relation.target_part
        extension = Path(relation.target_ref).suffix or ".png"
        filename = safe_filename(f"{docx_path.stem}_{index}", "image") + extension
        output_path = output_dir / filename
        output_path.write_bytes(part.blob)
        paths.append(str(output_path))
    return paths


def _store_pair(data: dict[str, str], label_text: str, value_text: str) -> None:
    label = label_text.strip().rstrip(":：").strip()
    field = _FIELD_MAP.get(label)
    if field is not None and value_text.strip():
        data[field] = value_text.strip()


def _parse_record_cell(data: dict[str, str], text: str) -> None:
    """解析合并单元格中的“情况记录 / 异常”多行文本。"""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    record_lines: list[str] = []
    abnormal_lines: list[str] = []
    mode: str | None = None

    for line in lines:
        if line.startswith("情况记录"):
            mode = "record"
            remainder = line[len("情况记录"):].lstrip(":：").strip()
            if remainder:
                record_lines.append(remainder)
        elif line.startswith("异常"):
            mode = "abnormal"
            remainder = line[len("异常"):].lstrip(":：").strip()
            if remainder:
                abnormal_lines.append(remainder)
        elif mode == "record":
            record_lines.append(line)
        elif mode == "abnormal":
            abnormal_lines.append(line)

    if record_lines:
        data["record"] = "\n".join(record_lines)
    if abnormal_lines:
        data["abnormal_status"] = " ".join(abnormal_lines)


def _merge_values(values: Iterable[str], separator: str = " / ") -> str:
    unique: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in unique:
            unique.append(text)
    return separator.join(unique)

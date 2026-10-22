"""报表 docx 渲染。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common.models import ReportData

DEFAULT_REPORT_TITLE = "农村生活污水设施巡查记录表"
ABNORMAL_COLOR = RGBColor(217, 48, 37)


def set_cell_margins(cell, top: int = 100, start: int = 100, bottom: int = 100, end: int = 100) -> None:
    """设置单元格边距（单位：twip）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_chinese_font(run, font_name: str = "微软雅黑") -> None:
    """同时设置西文与中文字体。"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def generate_report(
    output_path: str | Path,
    data: ReportData,
    title: str = DEFAULT_REPORT_TITLE,
) -> Path:
    """按统一模板生成一份报表，返回实际保存路径。"""
    output_path = Path(output_path)
    document = Document()

    _add_title(document, title)
    table = _create_table(document)
    _fill_basic_info(table, data)
    _fill_record(table, data)
    _fill_images(table, data.image_paths)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def _add_title(document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(title))
    run.bold = True
    run.font.size = Pt(22)
    set_chinese_font(run)


def _create_table(document):
    table = document.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    widths = [Cm(3), Cm(6), Cm(3), Cm(6)]
    for row_index, row in enumerate(table.rows):
        if row_index < 3:
            row.height = Cm(1)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        else:
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.height = Cm(3 if row_index == 3 else 8)

        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
    return table


def _fill_basic_info(table, data: ReportData) -> None:
    _fill_pair(table, 0, 0, "乡镇：", 1, data.town)
    _fill_pair(table, 0, 2, "日期：", 3, data.date)
    _fill_pair(table, 1, 0, "设施名称：", 1, data.facility_name)
    _fill_pair(table, 1, 2, "巡查人员：", 3, data.inspector)
    _fill_pair(table, 2, 0, "处理工艺：", 1, data.process)
    _fill_pair(table, 2, 2, "设施编号：", 3, data.facility_code)


def _fill_pair(table, row: int, label_column: int, label: str, value_column: int, value: object) -> None:
    label_cell = table.cell(row, label_column)
    label_run = label_cell.paragraphs[0].add_run(label)
    set_chinese_font(label_run)
    label_run.font.size = Pt(10)
    label_run.font.bold = True

    value_cell = table.cell(row, value_column)
    value_run = value_cell.paragraphs[0].add_run(str(value or ""))
    set_chinese_font(value_run)
    value_run.font.size = Pt(10)


def _fill_record(table, data: ReportData) -> None:
    record_cell = table.cell(3, 0).merge(table.cell(3, 3))
    paragraph = record_cell.paragraphs[0]

    label_run = paragraph.add_run("情况记录：\n")
    label_run.bold = True
    set_chinese_font(label_run)

    content_run = paragraph.add_run(str(data.record or ""))
    set_chinese_font(content_run)

    abnormal_label = paragraph.add_run("\n异常：")
    abnormal_label.bold = True
    abnormal_label.font.color.rgb = ABNORMAL_COLOR
    set_chinese_font(abnormal_label)

    abnormal_value = paragraph.add_run(str(data.abnormal_status or ""))
    abnormal_value.bold = True
    abnormal_value.font.color.rgb = ABNORMAL_COLOR
    set_chinese_font(abnormal_value)


def _fill_images(table, image_paths: list[str]) -> None:
    photo_cell = table.cell(4, 0).merge(table.cell(4, 3))
    paragraph = photo_cell.paragraphs[0]
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    added = 0
    for raw_path in image_paths:
        path = Path(raw_path)
        if not path.is_file():
            continue
        try:
            run = paragraph.add_run()
            run.add_picture(str(path), width=Cm(7))
        except Exception:
            continue

        added += 1
        if added % 2 == 0:
            paragraph.add_run("\n")
        else:
            paragraph.add_run("   ")

    if added == 0:
        placeholder = paragraph.add_run("\n\n           （暂无照片）")
        placeholder.bold = True
        set_chinese_font(placeholder)

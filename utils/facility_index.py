"""设施编号 Excel 的查询与缓存。"""

from __future__ import annotations

import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

FACILITY_FILE_NAME = "编号集合.xlsx"
REQUIRED_COLUMNS = ("乡镇", "自然村", "设施编码", "设施名称", "处理工艺")

LogFunc = Callable[..., None]


@dataclass(frozen=True)
class FacilityInfo:
    code: str
    name: str
    process: str


class FacilityIndex:
    """懒加载 Excel，避免每次解析消息都重新读盘。

    只依赖 openpyxl，不引入 pandas/numpy，打包后的体积小很多。
    """

    def __init__(self, file_path: str | Path | None = None, log_func: LogFunc | None = None):
        self._file_path = Path(file_path) if file_path else None
        self._log = log_func
        self._rows: list[dict[str, str]] | None = None
        self._load_attempted = False
        self._lock = threading.Lock()

    def set_logger(self, log_func: LogFunc) -> None:
        self._log = log_func

    def _candidate_paths(self) -> list[Path]:
        candidates: list[Path] = []
        if self._file_path is not None:
            candidates.append(self._file_path)
        for base in (Path.cwd(), Path(__file__).resolve().parent.parent):
            candidates.append(base / FACILITY_FILE_NAME)
            candidates.append(base / "utils" / FACILITY_FILE_NAME)

        result: list[Path] = []
        seen: set[Path] = set()
        for path in candidates:
            resolved = path.resolve()
            if resolved not in seen:
                seen.add(resolved)
                result.append(path)
        return result

    def _load(self) -> list[dict[str, str]] | None:
        with self._lock:
            if self._load_attempted:
                return self._rows
            self._load_attempted = True

            for path in self._candidate_paths():
                if not path.is_file():
                    continue
                try:
                    rows = _read_rows(path)
                except Exception as exc:
                    self._warn(f"读取设施编号文件失败 {path}: {exc}")
                    continue

                self._rows = rows
                self._file_path = path
                return rows

            self._warn("未找到可用的设施编号文件（编号集合.xlsx），编号相关报表会进入异常列表")
            return None

    def lookup(self, town: str | None, village: str | None) -> FacilityInfo | None:
        if not town or not village:
            return None

        rows = self._load()
        if rows is None:
            return None

        town_text = str(town).strip()
        village_text = str(village).strip()
        matches = [
            row
            for row in rows
            if row["乡镇"] == town_text and row["自然村"] == village_text
        ]
        if len(matches) != 1:
            return None

        row = matches[0]
        code = row["设施编码"]
        name = row["设施名称"]
        process = row["处理工艺"]
        if not code or not name or not process:
            return None
        return FacilityInfo(code=code, name=name, process=process)

    def _warn(self, message: str) -> None:
        if self._log is not None:
            self._log(message, "WARNING")


def _read_rows(path: Path) -> list[dict[str, str]]:
    """读取工作簿第一个工作表，返回按必需列组织的行。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        # 与 pandas.read_excel 的默认行为一致：读取第一个工作表
        sheet = workbook.worksheets[0]
        rows_iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iterator)
        except StopIteration:
            raise ValueError("工作簿为空") from None

        headers = [_cell_text(value) for value in header_row]
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"缺少列 {missing}")

        index = {column: headers.index(column) for column in REQUIRED_COLUMNS}
        rows: list[dict[str, str]] = []
        for values in rows_iterator:
            if values is None:
                continue
            row = {
                column: _cell_text(values[i]) if i < len(values) else ""
                for column, i in index.items()
            }
            if any(row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
